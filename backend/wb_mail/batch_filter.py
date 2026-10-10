"""wb_mail v6 批次过滤。

优化点：
- 同批次多行体时仅归入首个线体，避免重复发送
- 记录跳过的行及原因，便于排查"为什么某批次没发到"
- 试产线截断逻辑修正（v4 中 stop 赋值后紧接 break，使后续判断成为死代码）
- 增加 qty 统计用于邮件摘要
"""
from __future__ import annotations

import datetime
import re
from typing import Any

from openpyxl import load_workbook

from wb_mail import config as C

log = C.log


def _col_index(ws):
    """返回 (固定表头名->{名:列索引}, 日期列->列索引)。"""
    hdr: dict[str, int] = {}
    date_cols: dict[datetime.date, int] = {}
    for idx, cell in enumerate(ws[C.HEADER_ROW]):
        v = cell.value
        if v is None:
            continue
        name = str(v).strip()
        hdr.setdefault(name, idx)
        d = None
        if isinstance(v, datetime.datetime):
            d = v.date()
        elif isinstance(v, datetime.date):
            d = v
        if d is not None:
            date_cols[d] = idx

    need = [C.COL_LINE, C.COL_BATCH, C.COL_BOM]
    missing = [n for n in need if n not in hdr]
    if missing:
        raise KeyError(f"表头缺列: {missing}（当前表头: {list(hdr.keys())[:15]}）")
    return hdr, date_cols


def _num(v: Any) -> float | None:
    """把单元格解析为正数数量；非数字/空/占位符/布尔 返回 None。"""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if v > 0 else None
    s = str(v).strip().replace(",", "")
    if s in ("", "-", "—", "/", "无"):
        return None
    try:
        f = float(s)
    except ValueError:
        return None
    return f if f > 0 else None


def _has_qty(v: Any) -> bool:
    return _num(v) is not None


def _qty(v: Any) -> int | float:
    n = _num(v)
    if n is None:
        return 0
    return int(n) if n == int(n) else n


def _row_value(row: tuple, index: int | None):
    if index is None or index < 0 or index >= len(row):
        return None
    return row[index]


def _parse_plan_row(row: tuple, row_idx: int, hdr: dict[str, int],
                    date_col: int | None, seen: set[str],
                    batch_lines: dict[str, str], skip_log: list[str]):
    """解析单行计划；返回 (line, batch, qty)，无效行返回 None。"""
    line = str(_row_value(row, hdr[C.COL_LINE]) or "").strip()
    if C.PILOT_MARK in line:
        return "STOP"

    bom = str(_row_value(row, hdr[C.COL_BOM]) or "").strip()
    batch = str(_row_value(row, hdr[C.COL_BATCH]) or "").strip()
    raw_qty = _row_value(row, date_col)
    if not bom or not re.match(C.BOM_FILTER_RE, bom) or not batch:
        return None
    if date_col is not None and not _has_qty(raw_qty):
        return None
    if batch in seen:
        skip_log.append(
            f"行{row_idx}: {batch} 已分配至 {batch_lines.get(batch, '?')}"
        )
        return None
    if not line:
        skip_log.append(f"行{row_idx}: {batch} 线体为空，跳过")
        return None
    seen.add(batch)
    batch_lines[batch] = line
    return line, batch, (_qty(raw_qty) if date_col is not None else 1)


def extract_batches(path: str, plan_date: datetime.date | None = None) -> list[dict]:
    """返回 [{line, batches[], qty_by_batch{}, total_qty}]，
    仅含目标日期列有排产的 Z?U 批次。

    plan_date: datetime.date；None 则不按日期过滤（取全部，仅用于排障）。
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    seen_batches: set[str] = set()  # 全局去重：同批次只归入首个线体
    batch_lines: dict[str, str] = {}
    skip_log: list[str] = []

    try:
        ws = wb.active
        hdr, date_cols = _col_index(ws)

        dcol: int | None = None
        if plan_date is not None:
            if plan_date not in date_cols:
                log.info("计划表无 %s 对应日期列，跳过", plan_date)
                return []
            dcol = date_cols[plan_date]

        groups: dict[str, dict] = {}
        for row_idx, row in enumerate(ws.iter_rows(min_row=C.HEADER_ROW + 1,
                                                    values_only=True),
                                       start=C.HEADER_ROW + 1):
            row = row or (None,)
            parsed = _parse_plan_row(
                row, row_idx, hdr, dcol, seen_batches, batch_lines, skip_log
            )
            if parsed == "STOP":
                line = str(_row_value(row, hdr[C.COL_LINE]) or "").strip()
                log.debug("行%d 命中试产标记(%s)，截断", row_idx, line)
                break
            if parsed is None:
                continue
            line, batch, qty = parsed
            rec = groups.setdefault(line, {
                "line": line, "batches": [], "qty_by_batch": {}, "total_qty": 0})
            rec["batches"].append(batch)
            rec["qty_by_batch"][batch] = qty
            rec["total_qty"] += qty

        if skip_log:
            log.info("跳过 %d 行: %s", len(skip_log), "; ".join(skip_log[:10]))

    finally:
        wb.close()

    return list(groups.values())
